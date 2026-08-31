"""Тесты уведомлений: тихие часы, умное время, /mute, сверка выплаты, выгрузка.

Запуск: .venv/bin/python test_notify.py
"""
from __future__ import annotations

import csv
import io
import os
import tempfile
from datetime import date, datetime, time, timedelta
from decimal import Decimal

os.environ.setdefault("SHIFTBOT_DB", os.path.join(tempfile.mkdtemp(), "test.db"))

import config  # noqa: E402
import db  # noqa: E402
import domain  # noqa: E402
import exporting  # noqa: E402
import notify  # noqa: E402
import parsing  # noqa: E402
import reports  # noqa: E402

DEFAULT = notify.Prefs()                                   # тихие часы 23:00–07:00
NO_QUIET = notify.Prefs(quiet_from=time(0, 0), quiet_to=time(0, 0))
LONG_QUIET = notify.Prefs(quiet_from=time(20, 0), quiet_to=time(8, 0))


def at(y: int, m: int, d: int, hh: int, mm: int = 0) -> datetime:
    return datetime(y, m, d, hh, mm, tzinfo=config.TZ)


# --- Тихие часы -----------------------------------------------------------

def test_quiet_window():
    # окно через полночь: 23:00–07:00
    assert notify.in_quiet(DEFAULT, at(2026, 8, 28, 23, 30))
    assert notify.in_quiet(DEFAULT, at(2026, 8, 28, 3, 0))
    assert not notify.in_quiet(DEFAULT, at(2026, 8, 28, 7, 0))
    assert not notify.in_quiet(DEFAULT, at(2026, 8, 28, 20, 0))
    # выключенные тихие часы не срабатывают никогда
    assert not notify.in_quiet(NO_QUIET, at(2026, 8, 28, 3, 0))
    assert NO_QUIET.quiet_off and NO_QUIET.quiet_label == "выключены"
    assert DEFAULT.quiet_label == "23:00–07:00"
    # окно внутри одних суток тоже работает
    day_quiet = notify.Prefs(quiet_from=time(1, 0), quiet_to=time(6, 0))
    assert notify.in_quiet(day_quiet, at(2026, 8, 28, 2, 0))
    assert not notify.in_quiet(day_quiet, at(2026, 8, 28, 23, 0))


def test_quiet_end():
    assert notify.quiet_end(DEFAULT, at(2026, 8, 28, 23, 30)) == at(2026, 8, 29, 7, 0)
    assert notify.quiet_end(DEFAULT, at(2026, 8, 29, 2, 0)) == at(2026, 8, 29, 7, 0)
    # вне тихих часов момент не двигается
    assert notify.quiet_end(DEFAULT, at(2026, 8, 28, 12, 0)) == at(2026, 8, 28, 12, 0)


# --- Напоминание перед сменой ---------------------------------------------

def test_remind_at_morning_goes_to_evening_before():
    """1я смена в 06:00: «за 2 часа» — это 04:00, поэтому предупреждаем вечером."""
    monday = date(2026, 8, 31)
    assert notify.remind_at(DEFAULT, monday, 1) == at(2026, 8, 30, 21, 0)
    # дневная и ночная остаются на «за 2 часа» — там будить никого не надо
    assert notify.remind_at(DEFAULT, monday, 2) == at(2026, 8, 31, 12, 0)
    assert notify.remind_at(DEFAULT, monday, 3) == at(2026, 8, 31, 20, 0)


def test_remind_at_respects_evening_hour():
    monday = date(2026, 8, 31)
    prefs = notify.Prefs(evening_hour=19)
    assert notify.remind_at(prefs, monday, 1) == at(2026, 8, 30, 19, 0)


def test_remind_at_without_quiet_hours():
    """Без тихих часов работает старое поведение — ровно за 2 часа."""
    assert notify.remind_at(NO_QUIET, date(2026, 8, 31), 1) == at(2026, 8, 31, 4, 0)


def test_remind_at_falls_back_before_quiet_starts():
    """Если и вечер накануне попал в тихие часы, предупреждаем до их начала."""
    assert notify.remind_at(LONG_QUIET, date(2026, 8, 31), 1) == at(2026, 8, 30, 19, 59)
    # ночная смена начинается ровно в тихие часы — та же страховка
    assert notify.remind_at(LONG_QUIET, date(2026, 8, 31), 3) == at(2026, 8, 31, 19, 59)


def test_leave_at():
    """«Пора выходить» — фиксированные 40 минут, тихие часы тут не при чём."""
    assert notify.leave_at(date(2026, 8, 31), 1) == at(2026, 8, 31, 5, 20)
    assert notify.leave_at(date(2026, 8, 31), 3) == at(2026, 8, 31, 21, 20)


# --- Вопрос «сколько часов зачли» -----------------------------------------

def test_confirm_at_night_shift_asks_in_the_afternoon():
    """Ночная кончается в 06:00 — спрашиваем не в 06:15, а днём."""
    assert notify.confirm_at(DEFAULT, date(2026, 8, 31), 3) == at(2026, 9, 1, 14, 0)
    # дневные — как раньше, через 15 минут после конца
    assert notify.confirm_at(DEFAULT, date(2026, 8, 31), 1) == at(2026, 8, 31, 14, 15)
    assert notify.confirm_at(DEFAULT, date(2026, 8, 31), 2) == at(2026, 8, 31, 22, 15)


def test_confirm_at_moves_out_of_quiet_hours():
    """Тихие часы с 22:00: вопрос про 2ю смену уезжает на утро."""
    prefs = notify.Prefs(quiet_from=time(22, 0), quiet_to=time(7, 0))
    assert notify.confirm_at(prefs, date(2026, 8, 31), 2) == at(2026, 9, 1, 7, 0)


# --- /mute ----------------------------------------------------------------

def test_mute_until_end_of_day():
    assert notify.mute_until(DEFAULT, at(2026, 8, 31, 10, 0)) == at(2026, 9, 1, 0, 0)
    # поздним вечером «до конца дня» бессмысленно — глушим до утра
    assert notify.mute_until(DEFAULT, at(2026, 8, 31, 23, 30)) == at(2026, 9, 1, 7, 0)
    assert notify.mute_until_morning(NO_QUIET, at(2026, 8, 31, 23, 30)) == at(2026, 9, 1, 8, 0)


def test_muted_and_silent():
    prefs = notify.Prefs(muted_until=at(2026, 8, 31, 20, 0))
    assert notify.muted(prefs, at(2026, 8, 31, 12, 0))
    assert not notify.muted(prefs, at(2026, 8, 31, 20, 0))
    assert notify.silent(prefs, at(2026, 8, 31, 12, 0))
    # без /mute молчим только в тихие часы
    assert notify.silent(DEFAULT, at(2026, 8, 31, 3, 0))
    assert not notify.silent(DEFAULT, at(2026, 8, 31, 12, 0))


# --- Настройки в базе -----------------------------------------------------

def test_prefs_roundtrip():
    db.init()
    uid = 9001
    db.ensure_user(uid)
    prefs = db.prefs(uid)
    assert prefs.user_id == uid
    # по умолчанию всё включено, а тихие часы — как в config
    assert all(getattr(prefs, flag) for flag in db.PREF_FLAGS)
    assert prefs.quiet_from == config.QUIET_FROM and prefs.quiet_to == config.QUIET_TO
    assert prefs.evening_hour == config.EVENING_HOUR and prefs.muted_until is None

    db.set_flag(uid, "leave_ping", False)
    db.set_flag(uid, "monday_plan", False)
    db.set_quiet(uid, time(22, 0), time(6, 30))
    db.set_evening_hour(uid, 19)
    until = at(2026, 8, 31, 22, 0)
    db.set_muted_until(uid, until)

    prefs = db.prefs(uid)
    assert prefs.leave_ping is False and prefs.monday_plan is False
    assert prefs.weekly_ask is True                  # чужие настройки не задели
    assert prefs.quiet_from == time(22, 0) and prefs.quiet_to == time(6, 30)
    assert prefs.evening_hour == 19
    assert prefs.muted_until == until
    assert uid not in db.users_with("monday_plan")
    assert uid in db.users_with("period_news")

    # сброс к значениям из config
    db.set_quiet(uid, None, None)
    db.set_evening_hour(uid, None)
    db.set_muted_until(uid, None)
    prefs = db.prefs(uid)
    assert prefs.quiet_from == config.QUIET_FROM and prefs.evening_hour == config.EVENING_HOUR
    assert prefs.muted_until is None

    assert uid in db.all_prefs()
    try:
        db.set_flag(uid, "нет-такого", True)
    except ValueError:
        pass
    else:
        raise AssertionError("неизвестная настройка должна отвергаться")


# --- Выборки планировщика -------------------------------------------------

def test_due_reminders_use_evening_time():
    db.init()
    uid = 9002
    db.ensure_user(uid)
    monday = date(2026, 8, 31)
    db.add_shift(uid, monday, 1)          # 1я смена, напоминание вечером 30.08 в 21:00
    mine = lambda moment: [s for s in db.due_reminders(moment) if s.user_id == uid]

    assert mine(at(2026, 8, 30, 20, 0)) == []      # ещё рано
    ready = mine(at(2026, 8, 30, 21, 1))
    assert len(ready) == 1
    # в 04:00 напоминание уже отправлено бы; проверим, что до вечера его не было
    db.mark_reminded(ready[0].id)
    assert mine(at(2026, 8, 31, 4, 1)) == []

    # с выключенными тихими часами то же напоминание встаёт на 04:00
    db.set_quiet(uid, time(0, 0), time(0, 0))
    db.add_shift(uid, date(2026, 9, 1), 1)
    later = [s for s in db.due_reminders(at(2026, 9, 1, 4, 1)) if s.user_id == uid]
    assert len(later) == 1
    assert [s for s in db.due_reminders(at(2026, 8, 31, 21, 1)) if s.user_id == uid] == []
    db.set_quiet(uid, None, None)


def test_leave_reminders():
    db.init()
    uid = 9003
    db.ensure_user(uid)
    d = date(2026, 8, 31)
    db.add_shift(uid, d, 2)               # смена 14:00, «пора выходить» в 13:20
    mine = lambda moment: [s for s in db.due_leave_reminders(moment) if s.user_id == uid]

    assert mine(at(2026, 8, 31, 13, 0)) == []
    ready = mine(at(2026, 8, 31, 13, 25))
    assert len(ready) == 1
    assert mine(at(2026, 8, 31, 14, 1)) == []      # смена началась — поздно
    db.mark_leave_reminded(ready[0].id)
    assert mine(at(2026, 8, 31, 13, 25)) == []     # второй раз не шлём
    # обычное напоминание живёт своей жизнью
    assert [s for s in db.due_reminders(at(2026, 8, 31, 12, 1)) if s.user_id == uid]


def test_awaiting_ask_night_shift():
    db.init()
    uid = 9004
    db.ensure_user(uid)
    d = date(2026, 8, 31)
    db.add_shift(uid, d, 3)               # ночная 22:00–06:00
    mine = lambda moment: [s for s in db.shifts_awaiting_ask(moment) if s.user_id == uid]

    assert mine(at(2026, 9, 1, 6, 20)) == []       # человек только пришёл со смены
    ready = mine(at(2026, 9, 1, 14, 5))
    assert len(ready) == 1
    db.mark_confirm_asked(ready[0].id)
    assert mine(at(2026, 9, 1, 15, 0)) == []


def test_ignored_shifts_only_after_question():
    db.init()
    uid = 9005
    db.ensure_user(uid)
    db.add_shift(uid, date(2026, 8, 3), 1)
    db.add_shift(uid, date(2026, 8, 4), 1)
    shifts = db.shifts_in_range(uid, date(2026, 8, 3), date(2026, 8, 4))
    assert db.ignored_shifts(uid) == []            # ещё не спрашивали — не пингуем
    db.mark_confirm_asked(shifts[0].id)
    ignored = db.ignored_shifts(uid)
    assert [s.id for s in ignored] == [shifts[0].id]
    # ответили — пинг больше не нужен
    db.confirm_shift(shifts[0].id, uid, Decimal("8"))
    assert db.ignored_shifts(uid) == []


# --- Сверка выплаты -------------------------------------------------------

def test_payout_roundtrip():
    db.init()
    uid = 9006
    db.ensure_user(uid)
    assert db.get_payout(uid, 2026, 8) is None
    db.set_payout(uid, 2026, 8, Decimal("4520.35"))
    assert db.get_payout(uid, 2026, 8) == Decimal("4520.35")
    db.set_payout(uid, 2026, 8, Decimal("4600"))    # перезапись, а не второй ряд
    assert db.get_payout(uid, 2026, 8) == Decimal("4600")
    assert db.get_payout(uid, 2026, 9) is None      # чужой период не задет
    assert db.delete_payout(uid, 2026, 8) is True
    assert db.delete_payout(uid, 2026, 8) is False
    assert db.get_payout(uid, 2026, 8) is None


def test_payout_check_shows_difference():
    db.init()
    uid = 9007
    db.ensure_user(uid)
    db.add_shift(uid, date(2026, 8, 3), 1)
    db.add_shift(uid, date(2026, 8, 4), 1)
    period = reports.period_data(uid, 2026, 8)

    assert "не отмечена" in reports.payout_check(uid, 2026, 8)
    assert reports.payout_lines(uid, period) == []

    # пришло ровно столько, сколько посчитал бот
    db.set_payout(uid, 2026, 8, period.net)
    text = reports.payout_check(uid, 2026, 8)
    assert "Сходится" in text
    assert "сходится с расчётом" in " ".join(reports.payout_lines(uid, period))

    # пришло на 100 меньше
    db.set_payout(uid, 2026, 8, period.net - Decimal("100"))
    text = reports.payout_check(uid, 2026, 8)
    assert "Разница" in text and "меньше" in text
    assert domain.money(Decimal("100")) in text
    # и это видно прямо в /money
    assert any("Пришло" in line for line in reports.payout_lines(uid, period))
    assert domain.money(period.net) in reports.money_report(uid, 2026, 8)


# --- Тексты уведомлений и выгрузка ----------------------------------------

def test_period_and_payout_messages():
    db.init()
    uid = 9008
    db.ensure_user(uid)
    # пустой период — молчим, а не присылаем нули
    assert reports.period_close_message(uid, 2026, 8) is None
    assert reports.payout_message(uid, 2026, 8) is None

    db.add_shift(uid, date(2026, 8, 3), 1)
    db.add_penalty(uid, date(2026, 8, 5), "late")
    closed = reports.period_close_message(uid, 2026, 8)
    assert closed is not None
    assert domain.period_title(2026, 8) in closed
    assert "Выплата" in closed and f"{domain.payout_date(2026, 8):%d}" in closed

    payout = reports.payout_message(uid, 2026, 8)
    assert payout is not None and "выплата" in payout.lower()
    period = reports.period_data(uid, 2026, 8)
    assert domain.money(period.net) in payout


def test_unconfirmed_ping_text():
    db.init()
    uid = 9009
    db.ensure_user(uid)
    db.add_shift(uid, date(2026, 8, 3), 1)
    db.add_shift(uid, date(2026, 8, 4), 2)
    pending = db.unconfirmed_shifts(uid)
    text = reports.unconfirmed_ping(pending)
    assert "без подтверждения" in text
    assert "/confirm" in text
    assert domain.fmt_hours(config.SHIFT_HOURS) in text


def test_monday_plan_message():
    db.init()
    uid = 9010
    db.ensure_user(uid)
    real_today = domain.today
    monday = date(2026, 8, 31)
    next_monday = monday + timedelta(days=7)
    try:
        domain.today = lambda: monday
        text = reports.monday_plan_message(uid)
        assert text is not None
        assert domain.week_label(next_monday) in text
        # как только смены на следующей неделе записаны — молчим
        db.add_shift(uid, next_monday, 1)
        assert reports.monday_plan_message(uid) is None
    finally:
        domain.today = real_today


def test_settings_report():
    db.init()
    uid = 9011
    db.ensure_user(uid)
    text = reports.settings_report(uid)
    assert "Тихие часы" in text and "23:00–07:00" in text
    db.set_quiet(uid, time(0, 0), time(0, 0))
    db.set_flag(uid, "confirm_ping", False)
    text = reports.settings_report(uid)
    assert "Тихие часы выключены" in text
    assert "пинг про неподтверждённые" not in text


def test_export_xlsx():
    """Выгрузка — книга Excel: у каждого блока свой лист, числа остаются числами."""
    db.init()
    uid = 9012
    db.ensure_user(uid)
    db.add_shift(uid, date(2026, 8, 3), 1)
    db.add_shift(uid, date(2026, 8, 4), 3)
    db.add_penalty(uid, date(2026, 8, 5), "late")
    db.set_payout(uid, 2026, 8, Decimal("1000"))

    assert exporting.HAVE_XLSX, "openpyxl должен стоять — он в requirements.txt"
    assert exporting.export_name(2026, 8) == "smeny-2026-08.xlsx"
    raw = exporting.period_file(uid, 2026, 8)
    assert raw.startswith(b"PK")                    # xlsx это zip

    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(raw))
    assert book.sheetnames == ["Смены", "Штрафы", "Итоги"]

    smeny = book["Смены"]
    assert [c.value for c in smeny[1]][:6] == [
        "Дата", "День", "Смена", "Время", "Статус", "Часы",
    ]
    # даты — датами, часы и суммы — числами, иначе сверять невозможно
    assert smeny["A2"].value.date() == date(2026, 8, 3)
    assert smeny["F2"].value == 8 and isinstance(smeny["F2"].value, (int, float))
    assert smeny["H2"].value == 265.20
    assert smeny["G3"].value == 35.20                # ночная ставка
    # итог строкой-формулой: поправишь часы — пересчитается
    assert str(smeny["F4"].value).startswith("=SUM(")
    assert smeny.freeze_panes == "A2"

    fines = book["Штрафы"]
    assert fines["B2"].value == "Опоздание" and fines["C2"].value == 150.0

    totals = {r[0].value: r[1].value for r in book["Итоги"].iter_rows(min_row=2)}
    assert totals["Заработано"] == 546.80           # 265,20 + 281,60
    assert totals["Штрафы (1)"] == -150.0
    assert totals["На руки по расчёту"] == 396.80
    assert totals["Пришло фактически"] == 1000.0
    assert round(totals["Разница"], 2) == 603.20
    assert domain.period_title(2026, 8) in exporting.caption(uid, 2026, 8)


def test_export_csv_fallback_is_one_flat_table():
    """Без openpyxl отдаём CSV — но одной ровной таблицей, без склеенных секций."""
    db.init()
    uid = 9013
    db.ensure_user(uid)
    db.add_shift(uid, date(2026, 8, 3), 1)
    db.add_penalty(uid, date(2026, 8, 5), "late")

    raw = exporting.period_csv(uid, 2026, 8)
    assert raw.startswith(b"\xef\xbb\xbf")          # BOM, иначе Excel испортит кириллицу
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    # ровно 8 колонок в каждой строке — иначе редакторы разъезжаются
    assert {len(r) for r in rows if r} == {8}, {len(r) for r in rows if r}
    assert rows[0][0] == "Дата" and rows[0][5] == "Часы"
    assert rows[1][0] == "03.08.2026"


def test_period_text_fits_a_message():
    db.init()
    uid = 9014
    db.ensure_user(uid)
    for day in range(3, 12):
        db.add_shift(uid, date(2026, 8, day), 1 if day % 2 else 2)
    text = exporting.period_text(uid, 2026, 8)
    assert len(text) < 4096                          # лимит одного сообщения Telegram
    assert domain.period_title(2026, 8) in text
    assert "На руки" in text


def test_parse_amount():
    assert parsing.parse_amount("4520") == Decimal("4520")
    assert parsing.parse_amount(" 4 520,35 ") == Decimal("4520.35")
    assert parsing.parse_amount("4520.35 zł") == Decimal("4520.35")
    assert parsing.parse_amount("4520,3") == Decimal("4520.3")
    assert parsing.parse_amount("0") == Decimal("0")
    for junk in ("", "абракадабра", "4520,355", "-100", "4520 zł и ещё немного"):
        assert parsing.parse_amount(junk) is None, junk


def test_payout_anchor():
    # 17.09.2026 — четверг, это выплата за август
    assert domain.payout_anchor(date(2026, 9, 17)) == (2026, 8)
    assert domain.payout_anchor(date(2026, 9, 16)) is None
    # 17.01.2026 — суббота, выплата переехала на понедельник 19-го
    assert domain.payout_anchor(date(2026, 1, 19)) == (2025, 12)
    assert domain.payout_anchor(date(2026, 1, 17)) is None


if __name__ == "__main__":
    failed = 0
    for name, fn in sorted(list(globals().items())):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"ok   {name}")
            except AssertionError as exc:
                failed += 1
                print(f"FAIL {name}: {exc!r}")
            except Exception as exc:
                failed += 1
                print(f"ERR  {name}: {type(exc).__name__}: {exc}")
    print("\n" + ("все тесты прошли" if not failed else f"провалено: {failed}"))
    raise SystemExit(1 if failed else 0)
