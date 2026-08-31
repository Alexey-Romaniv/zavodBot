"""Выгрузка периода в таблицу — чтобы сверять с расчёткой завода.

Отдаём .xlsx, а не CSV: в CSV пришлось бы складывать в один файл блоки с разным
числом колонок (смены — восемь, штрафы — пять, итоги — две), и любой редактор
выравнивает такое по первой строке, разъезжаясь на остальных. В книге Excel
у каждого блока свой лист, даты остаются датами, а суммы — числами.

Итоги пишем посчитанными значениями, а не формулами: openpyxl формулу не
вычисляет и сохраняет без результата, поэтому Excel-то её пересчитает, а
Numbers и предпросмотр в Telegram покажут ноль.

Если openpyxl почему-то не установлен, отдаём CSV — но уже одной ровной
таблицей, без склеенных секций.
"""
from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal

import config
import db
import domain
import reports

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HAVE_XLSX = True
except ImportError:  # pragma: no cover — на машине без openpyxl
    HAVE_XLSX = False


DATE_FMT = "DD.MM.YYYY"
MONEY_FMT = '#,##0.00'
HOURS_FMT = '0.##'

HEAD_FILL = "FFE8EDF3"
TOTAL_FILL = "FFF6F3E8"


def status_label(shift: db.Shift) -> str:
    if shift.status == "absent":
        return "прогул"
    if shift.status == "cancelled":
        return "отменена"
    if shift.status == "done":
        return "отработано"
    if domain.shift_end(shift.work_date, shift.shift_num) <= domain.now():
        return "без подтверждения"
    return "запланировано"


def export_name(year: int, month: int) -> str:
    suffix = "xlsx" if HAVE_XLSX else "csv"
    return f"smeny-{year}-{month:02d}.{suffix}"


def period_file(user_id: int, year: int, month: int) -> bytes:
    if HAVE_XLSX:
        return period_xlsx(user_id, year, month)
    return period_csv(user_id, year, month)


# --- Excel ----------------------------------------------------------------

def _head(ws, row: int, titles: list[str]) -> None:
    thin = Side(style="thin", color="FFB8C4D4")
    for col, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)


def _widths(ws, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _total_row(ws, row: int, upto: int) -> None:
    """Подсветить итоговую строку — её видно сразу, не выискивая внизу."""
    for col in range(1, upto + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=TOTAL_FILL)


def _num(value: Decimal | int | float) -> float:
    return float(Decimal(str(value)))


def _shifts_sheet(ws, period: reports.Period) -> None:
    ws.title = "Смены"
    _head(ws, 1, [
        "Дата", "День", "Смена", "Время", "Статус",
        "Часы", f"Ставка, {config.CURRENCY}/ч", f"Сумма, {config.CURRENCY}",
    ])
    _widths(ws, [12, 6, 11, 14, 20, 8, 14, 14])
    ws.freeze_panes = "A2"

    shifts = sorted(
        period.counted + period.absent + period.cancelled,
        key=lambda s: (s.work_date, s.shift_num),
    )
    row = 2
    for s in shifts:
        counted = s.status not in ("absent", "cancelled")
        ws.cell(row=row, column=1, value=s.work_date).number_format = DATE_FMT
        ws.cell(row=row, column=2, value=domain.WEEKDAY_SHORT[s.work_date.weekday()])
        ws.cell(row=row, column=3, value=s.kind.title)
        ws.cell(row=row, column=4, value=s.kind.hours_label)
        ws.cell(row=row, column=5, value=status_label(s))
        ws.cell(row=row, column=6, value=_num(s.hours) if counted else 0).number_format = HOURS_FMT
        ws.cell(row=row, column=7, value=_num(s.kind.rate)).number_format = MONEY_FMT
        ws.cell(row=row, column=8, value=_num(s.pay) if counted else 0).number_format = MONEY_FMT
        row += 1

    if shifts:
        # Готовые числа, а не =SUM(): openpyxl формулы не вычисляет и пишет их без
        # посчитанного значения. Excel пересчитает при открытии, а Numbers,
        # предпросмотр в Telegram и на телефоне покажут в такой ячейке ноль.
        counted = [s for s in shifts if s.status not in ("absent", "cancelled")]
        ws.cell(row=row, column=5, value="Итого")
        ws.cell(
            row=row, column=6,
            value=_num(sum((s.hours for s in counted), Decimal(0))),
        ).number_format = HOURS_FMT
        ws.cell(
            row=row, column=8,
            value=_num(sum((s.pay for s in counted), Decimal(0))),
        ).number_format = MONEY_FMT
        _total_row(ws, row, 8)
    else:
        ws.cell(row=row, column=1, value="Смен за период не записано")


def _penalties_sheet(ws, period: reports.Period) -> None:
    ws.title = "Штрафы"
    _head(ws, 1, ["Дата", "Вид", f"Сумма, {config.CURRENCY}", "Основание по договору", "Заметка"])
    _widths(ws, [12, 28, 14, 48, 24])
    ws.freeze_panes = "A2"

    row = 2
    for pen in period.penalties:
        kind = domain.penalty(pen.kind)
        ws.cell(row=row, column=1, value=pen.at_date).number_format = DATE_FMT
        ws.cell(row=row, column=2, value=kind.title)
        ws.cell(row=row, column=3, value=_num(pen.amount)).number_format = MONEY_FMT
        ws.cell(row=row, column=4, value=kind.clause)
        ws.cell(row=row, column=5, value=pen.note or "")
        row += 1

    if period.penalties:
        ws.cell(row=row, column=2, value="Всего штрафов")
        ws.cell(
            row=row, column=3, value=_num(period.penalty_total)
        ).number_format = MONEY_FMT
        _total_row(ws, row, 5)
    else:
        ws.cell(row=row, column=1, value="Штрафов за период нет")


def _totals_sheet(ws, user_id: int, period: reports.Period) -> None:
    ws.title = "Итоги"
    _head(ws, 1, ["Показатель", "Значение"])
    _widths(ws, [34, 18])

    rows: list[tuple[str, object, str]] = [
        ("Период", domain.period_title(period.year, period.month), "text"),
        ("Начало периода", period.start, "date"),
        ("Конец периода", period.end, "date"),
        ("Дата выплаты", domain.payout_date(period.year, period.month), "date"),
        ("", "", "text"),
        ("Отработано часов (без плана)", _num(period.worked_hours), HOURS_FMT),
        (f"Заработано, {config.CURRENCY}", _num(period.earned), MONEY_FMT),
    ]
    if period.penalties:
        rows.append((
            f"Штрафы ({len(period.penalties)})", -_num(period.penalty_total), MONEY_FMT,
        ))
    if period.rate_cut:
        rows.append((
            f"Снижение ставки ({config.RATE_CUT} {config.CURRENCY}/ч)",
            -_num(period.rate_cut_amount), MONEY_FMT,
        ))
    rows.append((f"На руки по расчёту, {config.CURRENCY}", _num(period.net), MONEY_FMT))

    actual = db.get_payout(user_id, period.year, period.month)
    if actual is not None:
        rows.append(("Пришло фактически", _num(actual), MONEY_FMT))
        rows.append(("Разница", _num(actual - period.net), MONEY_FMT))
    rows += [
        ("", "", "text"),
        ("Ставка дневная", _num(config.DAY_RATE), MONEY_FMT),
        ("Ставка ночная", _num(config.NIGHT_RATE), MONEY_FMT),
    ]

    row = 2
    for label, value, fmt in rows:
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        if fmt == "date":
            cell.number_format = DATE_FMT
        elif fmt not in ("text",):
            cell.number_format = fmt
        if label.startswith(("На руки по расчёту", "Разница")):
            _total_row(ws, row, 2)
        row += 1


def period_xlsx(user_id: int, year: int, month: int) -> bytes:
    period = reports.period_data(user_id, year, month)
    book = Workbook()
    _shifts_sheet(book.active, period)
    _penalties_sheet(book.create_sheet(), period)
    _totals_sheet(book.create_sheet(), user_id, period)

    book.properties.title = f"Смены — {domain.period_title(year, month)}"
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


# --- CSV на случай отсутствия openpyxl ------------------------------------

def period_csv(user_id: int, year: int, month: int) -> bytes:
    """Одна ровная таблица: смены, потом штрафы и итоги теми же колонками.

    Никаких секций с другим числом столбцов — иначе редакторы разъезжаются.
    """
    period = reports.period_data(user_id, year, month)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    w.writerow(["Дата", "День", "Смена", "Время", "Статус", "Часы",
                f"Ставка, {config.CURRENCY}/ч", f"Сумма, {config.CURRENCY}"])

    def money(value: Decimal) -> str:
        return f"{Decimal(value).quantize(Decimal('0.01')):f}".replace(".", ",")

    def line(*cells: object) -> None:
        row = list(cells) + [""] * (8 - len(cells))
        w.writerow(row)

    for s in sorted(period.counted + period.absent + period.cancelled,
                    key=lambda x: (x.work_date, x.shift_num)):
        counted = s.status not in ("absent", "cancelled")
        w.writerow([
            f"{s.work_date:%d.%m.%Y}",
            domain.WEEKDAY_SHORT[s.work_date.weekday()],
            s.kind.title, s.kind.hours_label, status_label(s),
            money(s.hours) if counted else "0",
            money(s.kind.rate),
            money(s.pay) if counted else "0,00",
        ])

    line()
    line("Штрафы")
    for pen in period.penalties:
        kind = domain.penalty(pen.kind)
        line(f"{pen.at_date:%d.%m.%Y}", "", kind.title, kind.clause,
             pen.note or "", "", "", money(pen.amount))
    if not period.penalties:
        line("нет")

    line()
    line("Итоги")
    line("Отработано часов", "", "", "", "", money(period.worked_hours))
    line("Заработано", "", "", "", "", "", "", money(period.earned))
    if period.penalties:
        line("Штрафы", "", "", "", "", "", "", "-" + money(period.penalty_total))
    if period.rate_cut:
        line("Снижение ставки", "", "", "", "", "", "", "-" + money(period.rate_cut_amount))
    line("На руки по расчёту", "", "", "", "", "", "", money(period.net))
    actual = db.get_payout(user_id, year, month)
    if actual is not None:
        line("Пришло фактически", "", "", "", "", "", "", money(actual))
        line("Разница", "", "", "", "", "", "", money(actual - period.net))
    line("Выплата", "", "", "", "", "", "",
         f"{domain.payout_date(year, month):%d.%m.%Y}")

    # BOM: без него Excel читает файл как ANSI и портит кириллицу.
    return buf.getvalue().encode("utf-8-sig")


def caption(user_id: int, year: int, month: int) -> str:
    period = reports.period_data(user_id, year, month)
    n = len(period.counted)
    where = "Три листа: смены, штрафы, итоги." if HAVE_XLSX else "Открывается в Excel и Google Таблицах."
    return (
        f"📄 <b>{domain.period_title(year, month)}</b>"
        f" — {n} {domain.shifts_word(n)}, {domain.fmt_hours(period.worked_hours)}\n"
        f"На руки по расчёту: <b>{domain.money(period.net)}</b>\n"
        f"<i>{where}</i>"
    )


def period_text(user_id: int, year: int, month: int) -> str:
    """Та же выгрузка, но сообщением в чат — на телефоне это удобнее файла."""
    period = reports.period_data(user_id, year, month)
    lines = [
        f"📋 <b>{domain.period_title(year, month)}</b>",
        f"<i>{domain.fmt_date_long(period.start)} — {domain.fmt_date_long(period.end)}</i>",
        "",
    ]
    shifts = sorted(period.counted + period.absent + period.cancelled,
                    key=lambda s: (s.work_date, s.shift_num))
    if not shifts:
        lines.append("Смен за период не записано.")
    else:
        by_week: dict[date, list[db.Shift]] = {}
        for s in shifts:
            by_week.setdefault(domain.monday_of(s.work_date), []).append(s)
        for monday, items in sorted(by_week.items()):
            lines.append(f"<b>{domain.week_label(monday)}</b>")
            for s in items:
                mark = {"done": "✅", "absent": "🚫", "cancelled": "❌"}.get(s.status, "⏳")
                tail = "" if s.status in ("absent", "cancelled") else f" · {domain.money(s.pay)}"
                lines.append(
                    f"{mark} {domain.fmt_date(s.work_date)} {s.kind.title}"
                    f" · {domain.fmt_hours(s.hours)}{tail}"
                )
            lines.append("")
    lines.append(f"Отработано: <b>{domain.fmt_hours(period.worked_hours)}</b>")
    lines.append(f"Заработано: <b>{domain.money(period.earned)}</b>")
    if period.deductions:
        lines += reports.deduction_lines(period)
    lines.append(f"💵 На руки: <b>{domain.money(period.net)}</b>")
    lines += reports.payout_lines(user_id, period)
    lines.append(
        f"🗓 Выплата: <b>{domain.fmt_date_long(domain.payout_date(year, month))}</b>"
    )
    return "\n".join(lines)
